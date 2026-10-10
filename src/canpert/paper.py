"""Assemble the manuscript (Cell article format) -> .docx, then .pdf through MS Word."""

from __future__ import annotations

import json
import subprocess  # nosec - fixed local Word automation
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from canpert import results_text as R0
from canpert import results_v1 as R
from canpert.docxkit import BLUE, GREY, RED, Manuscript
from canpert.intro_text import INTRODUCTION
from canpert.methods_text import key_resources, method_details, quantification, reproducibility
from canpert.methods_formal import baselines_formal, downstream_formal
from canpert.theory_text import theory

TITLE = ("Cell line context reshapes genetic perturbation responses: set-based transfer of "
         "Perturb-seq effects across leukemia, hepatoblastoma, and non-transformed cells")


def _front(m: Manuscript, res: dict, figs: Path) -> None:
    d = m.doc
    p = d.add_paragraph()
    r = p.add_run("Article")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), GREY
    p = d.add_paragraph()
    r = p.add_run(TITLE)
    r.bold, r.font.size = True, Pt(17)
    p.paragraph_format.space_after = Pt(6)
    m.para("[Author names and affiliations to be added by the authors]", size=9, align="left")
    m.para("Correspondence: [corresponding author e-mail]", size=7.5, align="left", color=GREY)
    ga = figs / "GraphicalAbstract.png"
    if ga.exists():
        m.heading("Graphical abstract", 2)
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run().add_picture(str(ga), width=Inches(5.2))
    m.heading("Highlights", 2)
    m.bullets(R.highlights(res), size=8.5)
    m.heading("In brief", 2)
    m.para(R.in_brief(res), size=8.5)
    m.page_break()
    m.heading("Summary", 1)
    m.para(R.summary(res), bold=True, color=BLUE, size=9)
    m.para("**Keywords:** Perturb-seq; CRISPRi; cell line context; set transformer; transfer learning; "
           "calibrated metrics; p53; gene essentiality; DepMap; leukemia; hepatoblastoma", size=7.5)


def build_paper(rep: Path, pdf: bool = True) -> Path:
    """Build reports/CanPert_manuscript_V.1.0.docx (+ .pdf)."""
    res = json.loads((rep / "results.json").read_text())
    figs = rep / "figures"
    m = Manuscript(rep / "_eq")
    _front(m, res, figs)
    m.columns(2)
    m.heading("Introduction", 1)
    for t in INTRODUCTION:
        m.para(t, indent=True)
    m.heading("Results", 1)
    R.results(m, res, figs)
    m.columns(2, new_page=True)
    m.heading("Discussion", 1)
    for t in R.discussion(res):
        m.para(t, indent=True)
    m.heading("Limitations of the study", 2)
    for t in R.limitations(res):
        m.para(t, indent=True)
    m.heading("Resource availability", 1)
    m.heading("Lead contact", 3)
    m.para("Requests for further information should be directed to the lead contact [to be added].")
    m.heading("Materials availability", 3)
    m.para("This study did not generate new materials or new experimental data.")
    m.heading("Data and code availability", 3)
    m.bullets([
        "All data analysed are public: Replogle et al. 2022 and Nadig et al. 2025 CRISPRi Perturb-seq "
        "(Arc Institute knockdown-filtered concatenation on Hugging Face), DepMap 24Q4, STRING v12 and "
        "Reactome; identifiers are listed in the key resources table.",
        "All code (data ingestion, CASTER, baselines, PertBench, statistics, figures and this manuscript) "
        "is in the canpert repository [URL/DOI to be added on deposition]. Every number in the paper is "
        "generated from reports/results.json; every figure panel ships the CSV it is drawn from.",
        "The pre-specified protocol and its freeze point are in RESEARCH_PROTOCOL.md (git tag protocol-v1.0).",
    ])
    m.heading("Acknowledgments", 1)
    m.para("We thank the authors of the Replogle et al. and Nadig et al. Perturb-seq studies, the Arc "
           "Institute for releasing the processed data, and the DepMap, STRING and Reactome teams. "
           "[Funding to be added.]")
    m.heading("Author contributions", 1)
    m.para("[To be completed by the authors.]")
    m.heading("Declaration of interests", 1)
    m.para("The authors declare no competing interests. [To be confirmed by the authors.]")
    m.heading("Declaration of generative AI and AI-assisted technologies", 1)
    m.para("The authors used Claude (Anthropic), an AI assistant, as a coding and writing tool. It wrote the "
           "analysis code (data access, CASTER, baselines, evaluation and statistics), ran the computational "
           "experiments on the authors' workstation, produced the figures, and drafted and formatted the "
           "manuscript and its revisions, under the authors' direction. The study design, the pre-specified "
           "protocol and its amendments, and the interpretation were reviewed by the authors, who checked the "
           "results against the code and take full responsibility for the content.")
    m.heading("Supplemental information", 1)
    m.para("Tables S1-S11 and Figures S1-S15 with their legends follow the STAR Methods.")
    m.heading("References", 1)
    m.references_placeholder()
    # STAR methods
    m.columns(1, new_page=True)
    p = m.doc.add_paragraph()
    r = p.add_run("STAR★METHODS")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), RED
    key_resources(m)
    m.columns(2)
    method_details(m, res)
    baselines_formal(m)
    downstream_formal(m)
    theory(m)
    reproducibility(m)
    quantification(m)
    R0.supplementary(m, res, figs)
    m.fill_references()
    out = m.save(rep / "CanPert_manuscript_V.1.0.docx")
    if pdf:
        out = out.resolve()  # Word COM needs absolute paths
        target = out.with_suffix(".pdf")
        script = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
                  f"$d = $w.Documents.Open('{out}'); $d.SaveAs2('{target}', 17); $d.Close(); $w.Quit()")
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)  # noqa: S603, S607
    return out
