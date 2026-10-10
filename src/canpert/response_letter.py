"""Point-by-point response to the V.0.0 review (CanPert_review_and_todo.docx), numbers from results.json."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from canpert.results_text import f, pval

RED = RGBColor(0xA8, 0x1F, 0x2D)
LN = {"k562": "K562", "hepg2": "HepG2", "jurkat": "Jurkat", "rpe1": "RPE1"}


def _shade(cell, hex_: str) -> None:  # noqa: ANN001
    tcpr = cell._tc.get_or_add_tcPr()  # noqa: SLF001
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tcpr.append(sh)


def _table(doc: Document, header: list[str], rows: list[list[str]], widths: list[float]) -> None:
    from docx.shared import Inches

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(8)
        _shade(c, "F2E6E8")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(8)
    for row in t.rows:
        for c, w in zip(row.cells, widths, strict=True):
            c.width = Inches(w)
    doc.add_paragraph()


def build(rep: Path) -> Path:
    res = json.loads((rep / "results.json").read_text())
    rv = res.get("review", {})
    ex = rv.get("extra", {})
    ctr = rv.get("controls", {})
    cal = rv.get("calibration", {})
    seeds = rv.get("seeds", {})
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9.5)
    p = doc.add_paragraph()
    r = p.add_run("CanPert V.1.0: response to the V.0.0 review")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), RED
    doc.add_paragraph("Response to the review and to-do list of 8 October 2026. Every item was addressed. Numbers are "
                      "taken from reports/results.json of V.1.0. Additions made in this round are labelled "
                      "exploratory in the manuscript and recorded as AMENDMENT-003 in RESEARCH_PROTOCOL.md.")

    def head(t: str) -> None:
        h = doc.add_paragraph()
        rr = h.add_run(t)
        rr.bold, rr.font.size, rr.font.color.rgb = True, Pt(12), RED

    # ---------------------------------------------------------------- headline
    head("What changed in the conclusions")
    ca, pm, li, lo = (seeds.get("all_lines", {}).get(m, {}).get("pearson_delta", {}) for m in
                      ("CASTER", "PerturbMean", "Linear", "LineOffset"))
    cca, cpm, cli = (seeds.get("cancer_only", {}).get(m, {}).get("pearson_delta", {}) for m in ("CASTER", "PerturbMean", "Linear"))
    bullets = [
        f"CASTER vs Linear is a tie and is now stated as such everywhere (title, highlights, summary, Table 1). Over "
        f"{len(seeds.get('seeds', []))} seeds: CASTER {f(ca.get('mean'), 3)} ± {f(ca.get('sd'), 3)}, Linear {f(li.get('mean'), 3)} ± "
        f"{f(li.get('sd'), 3)}, PerturbMean {f(pm.get('mean'), 3)} ± {f(pm.get('sd'), 3)}, LineOffset {f(lo.get('mean'), 3)} "
        f"(Pearson delta, mean over lines). Cancer lines only: CASTER {f(cca.get('mean'))}, Linear {f(cli.get('mean'))}, "
        f"PerturbMean {f(cpm.get('mean'))}.",
        "Zero-shot is reported as a clear loss (" + f"{rv.get('t2_worse_panels', {}).get('worse')} of "
        f"{rv.get('t2_worse_panels', {}).get('total')} panels worse; H2 significant in the opposite direction).",
        "The DepMap link is reframed as general essentiality; the line-specific test (H4b) failing is stated plainly.",
        "HepG2 TP53 status is stated, and a p53 analysis shows shared p53 activity explains part of the HepG2-RPE1 result.",
        "Title changed to 'Cell line context reshapes genetic perturbation responses: ...'; 'pre-registered' changed to "
        "'pre-specified'; 'cancer target prioritisation' and 'benchmark for virtual-cell models' removed.",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    # ---------------------------------------------------------------- main issues
    head("Main issues (1-9)")
    cr = pd.DataFrame(ex.get("depmap_controls", {}).get("cross", []))
    ceil = pd.DataFrame(cal.get("ceiling", []))
    ceil1 = ceil[ceil.task == "t1"].set_index("target").ceiling if len(ceil) else pd.Series(dtype=float)
    lo_t = ctr.get("LineOffset", {}).get("pearson_delta", {})
    io_t = ctr.get("idonly", {}).get("pearson_delta", {})
    wc_t = ctr.get("wrongctrl", {}).get("pearson_delta", {})
    p53 = ex.get("p53", {})
    rows = [
        ["1", "CASTER ties Linear; Linear missing from figures",
         "Tie stated in title/highlights/summary/Table 1 and Results. Linear added to Figures 1G, 4C, 4D, 5D, S8, S10 "
         "(and was already in 2, 4A-B, 4E, 5C, 5E-F). Per-seed tests vs Linear reported."],
        ["2", "Gain may be one shift per line",
         f"New LineOffset baseline (PerturbMean + per-line offset): {f(lo_t.get('mean_b'))} vs CASTER {f(lo_t.get('mean_a'))}, "
         f"p = {pval(lo_t.get('p_two_sided'))}. Line-identity-only CASTER (no control cells): {f(io_t.get('mean_b'))} vs "
         f"{f(io_t.get('mean_a'))}, p = {pval(io_t.get('p_two_sided'))}. New section + Figure 6E. Pearson delta vs perturbed mean "
         "(top 100 DEGs) added as a metric that removes the shared shift."],
        ["3", "Most gain in RPE1", "Cancer-only averages reported for every seed (text, Figure 6D); RPE1 dependence stated."],
        ["4", "No noise ceiling", "Split-half noise ceilings per line (median): " + ", ".join(
            f"{LN[n]} {f(ceil1.get(n))}" for n in LN if n in ceil1) + ". Figure 6A; technical/interpolated duplicates as "
         "positive-control rows (Figure 6C, Table S10). Nadig 2025 checked: it reports r = 0.62 for transcriptome-wide "
         "impact *magnitude* across lines, not response-profile correlation; described accurately."],
        ["5", "Zero-shot is a clear loss", f"Reworded; {rv.get('t2_worse_panels', {}).get('worse')}/"
         f"{rv.get('t2_worse_panels', {}).get('total')} panels worse; two-sided p values shown in Table 1."],
        ["6", "One seed, one split", f"Seeds {seeds.get('seeds')} (each a new split and initialisation); mean +- SD, "
         "per-seed tests, cancer-only averages."],
        ["7", "HepG2 is TP53 wild type", "Stated; p53 target score analysis (Figure S13): induced in " + ", ".join(
            f"{LN[n]} {f(100 * p53.get('per_line', {}).get(n, {}).get('frac_induced', np.nan), 0)}%" for n in LN)
         + f"; HepG2-RPE1 p53 rank correlation {f(p53.get('pair_spearman_p53_shift', {}).get('hepg2-rpe1'))}. "
         "Shared p53 explains part, not all, of the similarity."],
        ["8", "DepMap link is general essentiality", "Cross-line control (Figure S14): response vs own and other lines' "
         "Chronos " + (", ".join(f"{LN[n]} own {f(cr[(cr.response_line == n) & (cr.chronos_line == n)].rho.iloc[0])}"
                                 for n in ('k562', 'hepg2', 'jurkat')) if len(cr) else "")
         + ". Reframed as general essentiality; H5 moved out of Table 1 and labelled exploratory."],
        ["9", "Weak metrics", "Added WMSE, weighted R2 (vs control and vs perturbed mean), top-100 Pearson delta vs "
         "perturbed mean, MSE, NIR, PDS, DRF heat map for our setup (Figure 6B, including the unseen-line task), "
         "following the Miller et al. code. Old metrics kept; new ones labelled as added later."],
    ]
    _table(doc, ["#", "Issue", "Response"], rows, [0.3, 1.6, 4.6])

    # ---------------------------------------------------------------- must-do experiments
    head("Must-do experiments (1-9)")
    gr = rv.get("grouped", {}).get("ribosome_gain_grouped", {})
    ab = pd.DataFrame(rv.get("ablation20", {}).get("means", []))
    abmu = ab.groupby("method").pearson_delta.mean() if len(ab) else pd.Series(dtype=float)
    rows = [
        ["1", "PerturbMean + line offset; line-ID-only model", "Done (LineOffset in every table; line-identity-only "
         "CASTER on all four lines). See issue 2."],
        ["2", "Noise ceiling", "Done (split-half + Spearman-Brown; technical and interpolated duplicates). 32 cells "
         "per knockdown were available, so halves have 16 cells; ceilings are reported for the full-cell truth."],
        ["3", "3+ seeds, splits, per line, cancer-only", "Done (seeds 0, 1, 2)."],
        ["4", "Wrong-line control cells", f"Done: {f(wc_t.get('mean_b'))} vs {f(wc_t.get('mean_a'))} with correct cells, "
         f"p = {pval(wc_t.get('p_two_sided'))} (seeds 1-2, all lines)."],
        ["5", "p53 check", "Done (Figure S13). CDKN1A and GDF15 are not among the measured genes; 14 consensus "
         "targets used (Fischer 2017)."],
        ["6", "DepMap control", "Done (Figure S14)."],
        ["7", "Ribosome claim: hold out complexes + enrichment", "Enrichment adjusted for response strength in every "
         "line (all held-out knockdowns); sibling-leakage check; complex-grouped splits (Louvain, STRING >= 950). "
         + ("Grouped-split median gain for ribosome knockdowns: " + ", ".join(
             f"{LN[n]} {f(gr[n]['median_gain_ribosome'], 3)}" for n in gr) if gr else "")],
        ["8", "Ablations at 20 epochs with seeds", "Done; full model " + f(abmu.get("S32"), 3) + ", PerturbMean "
         + f(abmu.get("PerturbMean"), 3) + ", Linear " + f(abmu.get("Linear"), 3) + " (K562)."],
        ["9", "Run a published model (State)", "Not run (CPU-only workstation). Instead the 'benchmark for virtual-cell "
         "models' wording was removed and the gap is stated in Discussion and Limitations."],
    ]
    _table(doc, ["#", "Experiment", "Status"], rows, [0.3, 1.9, 4.3])

    # ---------------------------------------------------------------- text fixes
    head("Fixes in the text, numbers and figures")
    rows = [
        ["'starts at the strongest simple baseline'", "Corrected to 'starts at the PerturbMean transfer baseline'; text "
         "notes ContextMean and Linear beat PerturbMean in some lines."],
        ["Refs 44, 45 (Cas9 cutting)", "Replaced by p53 / nucleolar-stress references (Fischer 2017; Golomb 2014; "
         "Pelava 2016; Lindstrom 2022)."],
        ["Simulation: Linear equals CASTER", "Reported in text; simulation design added to STAR Methods."],
        ["'no FiLM ... variant better' confusing", "Ablation text regenerated from the 20-epoch, two-seed results."],
        ["0.64/0.63 vs Figure 2E 0.65/0.53", "Text now uses the same pooled values as Figure 2E."],
        ["MAE test flipped", "Confirmed in code (lower-is-better metrics were tested as higher-is-better). Fixed; MAE "
         "now favours CASTER."],
        ["'indistinguishable from PerturbMean'", "Now: close on Pearson delta, worse in "
         f"{rv.get('transfer_worse_counts', {}).get('discrimination')}/12 directions on discrimination and "
         f"{rv.get('transfer_worse_counts', {}).get('de_overlap')}/12 on DE overlap."],
        ["Table 1 H2 / H5", "H2 shows two-sided p values of the opposite-direction tests; H5 removed from Table 1."],
        ["Fig 1G PerturbMean line", "Added (and Linear)."],
        ["Fig 2A, 3E '30%'", "Now '25% + 5%'."],
        ["Fig 4C, 4D, 5D, S8, S10", "Linear (and NoChange in 4D) added; S10 shows best and worst cases per line."],
        ["'same single-cell DE pipeline'", "Corrected: same pseudobulk pipeline incl. the amended predicted-side test."],
        ["Refs 17, 18 (nucleolar stress)", "Replaced by nucleolar-stress references."],
        ["'Tables S1-S8'; uncited items", "Now 'Tables S1-S11, Figures S1-S15'; every supplementary item is cited."],
        ["Placeholders, AI statement", "AI statement rewritten as what was done. Author list, contact, funding and "
         "code DOI remain placeholders: these need the authors' details."],
        ["'same library and protocol'", "Softened: same sgRNA library; screens differ in effector, time point, depth."],
        ["'about 35-50 min'", f"Now {f(rv.get('minutes_range', [np.nan, np.nan])[0], 0)}-"
         f"{f(rv.get('minutes_range', [np.nan, np.nan])[1], 0)} min (Table S9)."],
        ["'Both versions of every affected test'", "Table S1 now lists the pre-specified Wilcoxon version of every "
         "DE-affected test (DE AUPRC, DE overlap, fold-change Spearman, direction) for all baselines."],
        ["Table S8 odd labels, p values", "Terms restricted to 10-500 genes; BH-adjusted over terms; unenriched "
         "programs labelled '(none)'."],
        ["MAE undefined; 2,023 vs 1,850", "MAE/MSE defined. 1,850 of the 2,023 targets are among the 6,546 measured "
         "genes and were excluded from the gene space; the other 173 are not measured. Explained in Results."],
    ]
    _table(doc, ["Item", "Fix"], rows, [2.0, 4.5])

    # ---------------------------------------------------------------- wording
    head("Wording softened")
    rows = [["Title", "'Cell line context reshapes ...'"],
            ["Highlight 2", "'CASTER corrects transferred responses when part of the target line is screened'"],
            ["Highlight 3", "'... outperforms mean baselines and matches a linear model; it fails in unseen lines'"],
            ["Highlight 4", "'... tracks general gene essentiality, not line-specific dependency'"],
            ["Summary", "'In lines with no screening data, CASTER is worse than the perturbation mean.'"],
            ["'Unperturbed transcriptomes carry little information'", "'This model, trained on three lines, did not find ...'"],
            ["'Pre-registered'", "'Pre-specified' throughout (protocol frozen by git tag, not publicly time-stamped)."],
            ["'cancer target prioritisation'", "Removed."]]
    _table(doc, ["Where", "Now"], rows, [2.0, 4.5])

    # ---------------------------------------------------------------- double-check
    head("Double-check items")
    nc = rep / "tables" / "noise_cancel_check.csv"
    ncs = ""
    if nc.exists():
        d = pd.read_csv(nc)
        nf = rv.get("noise_free", {}).get("all_lines", {})
        def gg(mth: str) -> str:
            return f(nf.get(f"{mth}_clean", {}).get("pearson_delta", {}).get("mean"))
        ncs = ("Checked, and the answer is no, but the check found a real issue. Removing the sampling noise of the "
               "input control cells raised every method by a similar margin, CASTER included (noise-free Pearson "
               f"delta: CASTER {gg('CASTER')}, Linear {gg('Linear')}, LineOffset {gg('LineOffset')}, PerturbMean "
               f"{gg('PerturbMean')}; as rendered: {f(d.CASTER.mean())}, {f(d.Linear.mean())}, "
               f"{f(d.LineOffset.mean()) if 'LineOffset' in d else 'n/a'}, {f(d.PerturbMean.mean())} for seed 0). So CASTER "
               "does not gain by cancelling noise; the pre-specified rendering understated all absolute accuracies. "
               "The ranking is unchanged. New Results paragraph and Figure 6G.")
    ag = rep / "tables" / "similarity_allgenes.csv"
    ags = ""
    if ag.exists():
        d = pd.read_csv(ag)
        d["kind"] = np.where((d.a == "rpe1") | (d.b == "rpe1"), "with RPE1", "cancer-cancer")
        g = d.groupby("kind").pearson_allgenes.median()
        ags = (f"On all 6,546 measured genes, median cross-line similarity was {f(g.get('cancer-cancer'))} (cancer-cancer) "
               f"and {f(g.get('with RPE1'))} (with RPE1), versus 0.31 and 0.38 on the 2,000 selected genes.")
    rows = [["Nadig 2025 numbers", "Checked: 0.62 is the average cross-line correlation of transcriptome-wide impact "
             "(magnitude), not of response profiles; RPE1 most sensitive. Text corrected accordingly."],
            ["MAE direction", "Confirmed flipped in code; fixed."],
            ["CASTER cancelling control noise?", ncs or "See Table S11."],
            ["Gene choice favours lineage genes?", ags or "See Table S11."],
            ["DRF numbers read off a figure", "We now compute DRF ourselves on our splits (Figure 6B)."]]
    _table(doc, ["Item", "Result"], rows, [2.0, 4.5])
    out = rep / "CanPert_response_to_review_V.1.0.docx"
    doc.save(str(out))
    return out
