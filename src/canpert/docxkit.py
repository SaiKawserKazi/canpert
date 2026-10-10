"""Small python-docx toolkit for a Cell-style manuscript.

Features: Arial typography, coloured section headings, two-column body sections,
single-column figure pages, typeset equations (matplotlib mathtext -> 600 dpi PNG),
numbered citations ([@key] -> superscript numbers in order of first citation), tables.
"""

from __future__ import annotations

import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from canpert.references import REFS

RED = RGBColor(0xA8, 0x1F, 0x2D)  # Cell section-heading red
BLUE = RGBColor(0x00, 0x7D, 0xC5)  # Cell summary blue
GREY = RGBColor(0x55, 0x55, 0x55)
BODY_PT = 9.0


class Manuscript:
    """Stateful document builder."""

    def __init__(self, eq_dir: Path) -> None:
        self.doc = Document()
        self.eq_dir = eq_dir
        eq_dir.mkdir(parents=True, exist_ok=True)
        self.cite_order: list[str] = []
        self.n_eq = 0
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(BODY_PT)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.08
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.7))
        sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.7)
        self._footer(sec)

    # ------------------------------------------------------------------ layout
    def _footer(self, sec) -> None:  # noqa: ANN001
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
                run._r.append(el)  # noqa: SLF001
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                run._r.append(it)  # noqa: SLF001
        run.font.size = Pt(7)
        run.font.color.rgb = GREY

    def _cols(self, sec, n: int) -> None:  # noqa: ANN001
        cols = sec._sectPr.xpath("./w:cols")  # noqa: SLF001
        el = cols[0] if cols else OxmlElement("w:cols")
        el.set(qn("w:num"), str(n))
        el.set(qn("w:space"), "360")
        if not cols:
            sec._sectPr.append(el)  # noqa: SLF001

    def columns(self, n: int, new_page: bool = False) -> None:
        """Start a new section with n text columns."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
        self._cols(sec, n)

    def page_break(self) -> None:
        """Hard page break."""
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ text
    def _cite(self, text: str) -> list[tuple[str, bool]]:
        """Split text into (chunk, is_superscript) after resolving [@k1,k2] markers."""
        out: list[tuple[str, bool]] = []
        pos = 0
        for m in re.finditer(r"\[@([^\]]+)\]", text):
            out.append((text[pos:m.start()], False))
            nums = []
            for key in [k.strip() for k in m.group(1).split(",")]:
                if key not in REFS:
                    raise KeyError(f"unknown reference {key}")
                if key not in self.cite_order:
                    self.cite_order.append(key)
                nums.append(self.cite_order.index(key) + 1)
            out.append((_compress(sorted(set(nums))), True))
            pos = m.end()
        out.append((text[pos:], False))
        return out

    def para(self, text: str, *, bold: bool = False, italic: bool = False, size: float | None = None,
             color: RGBColor | None = None, align: str = "justify", space_after: float = 3,
             indent: bool = False) -> None:
        """Body paragraph; supports [@cite] markers and **bold** / *italic* spans."""
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                       "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(space_after)
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.15)
        text = text.replace("= < ", "< ").replace("p = <", "p <")
        self._runs(p, text, bold=bold, italic=italic, size=size, color=color)

    def _runs(self, p, text: str, *, bold: bool = False, italic: bool = False,  # noqa: ANN001
              size: float | None = None, color: RGBColor | None = None) -> None:
        """Add runs for text with citations, **bold**/*italic* spans and inline $math$."""
        for chunk, sup in self._cite(text):
            for k, seg in enumerate(re.split(r"\$([^$]+)\$", chunk)):
                if k % 2 == 1:
                    self._inline_math(p, seg, size or BODY_PT)
                    continue
                for piece, b2, i2 in _md_spans(seg):
                    r = p.add_run(piece)
                    r.bold = bold or b2
                    r.italic = italic or i2
                    r.font.superscript = sup
                    if size:
                        r.font.size = Pt(size)
                    if color is not None:
                        r.font.color.rgb = color

    def _inline_math(self, p, latex: str, pt: float) -> None:  # noqa: ANN001
        """Typeset inline math at the body font size and embed it in the line."""
        self.n_inline = getattr(self, "n_inline", 0) + 1
        path = self.eq_dir / f"inline{self.n_inline:04d}.png"
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=pt, family="serif")
        fig.savefig(path, dpi=600, bbox_inches="tight", pad_inches=0.005, facecolor="white")
        plt.close(fig)
        from PIL import Image

        with Image.open(path) as im:
            w_px, h_px = im.size
        run = p.add_run()
        run.add_picture(str(path), width=Inches(w_px / 600), height=Inches(h_px / 600))

    def heading(self, text: str, level: int = 1) -> None:
        """Cell-style headings: L1 = red caps, L2 = red bold, L3 = red bold italic."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text.upper() if level == 1 else text)
        r.bold = True
        r.italic = level == 3
        r.font.size = Pt(9 if level == 1 else 8.5)
        r.font.color.rgb = RED

    def bullets(self, items: list[str], size: float | None = None) -> None:
        """Bulleted list."""
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1.5)
            self._runs(p, it, size=size)

    def equation(self, latex: str, width_in: float | None = None, fontsize: int = 11) -> None:
        """Centered, typeset equation with a right-aligned number."""
        self.n_eq += 1
        path = self.eq_dir / f"eq{self.n_eq:03d}.png"
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=fontsize, family="serif")
        fig.savefig(path, dpi=600, bbox_inches="tight", pad_inches=0.02, transparent=False,
                    facecolor="white")
        plt.close(fig)
        from PIL import Image

        with Image.open(path) as im:
            w_px, h_px = im.size
        w_in = w_px / 600 * 0.85 if width_in is None else width_in
        w_in = min(w_in, 2.8)
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Inches(2.95), Inches(0.35)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run().add_picture(str(path), width=Inches(w_in))
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p1.add_run(f"({self.n_eq})")
        r.font.size = Pt(7.5)
        _ = h_px

    def figure(self, path: Path, title: str, legend: list[str], width_in: float = 7.0) -> None:
        """Full-width figure page: image, bold red title, legend paragraphs."""
        self.columns(1, new_page=True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width_in))
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RED
        for line in legend:
            self.para(line, size=7, space_after=1.5)

    def table(self, header: list[str], rows: list[list[str]], widths: list[float] | None = None,
              size: float = 6.8) -> None:
        """Simple ruled table."""
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RED
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for c, w in zip(row.cells, widths, strict=True):
                    c.width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def references_placeholder(self) -> None:
        """Reserve the position of the reference list (filled once all citations are known)."""
        self._ref_anchor = self.doc.add_paragraph()

    def fill_references(self) -> None:
        """Insert the numbered list at the reserved anchor."""
        anchor = self._ref_anchor
        for i, key in enumerate(self.cite_order, 1):
            p = anchor.insert_paragraph_before()
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            r = p.add_run(f"{i}. {REFS[key]}")
            r.font.size = Pt(6.8)

    def references(self) -> None:
        """Numbered reference list in order of first citation."""
        for i, key in enumerate(self.cite_order, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            r = p.add_run(f"{i}. {REFS[key]}")
            r.font.size = Pt(6.8)

    def save(self, path: Path) -> Path:
        """Write the .docx."""
        self.doc.save(str(path))
        return path


def _compress(nums: list[int]) -> str:
    """[1,2,3,5] -> '1-3,5' (Cell style uses en dash)."""
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1])))
        i = j + 1
    return ",".join(out)


def _md_spans(text: str) -> list[tuple[str, bool, bool]]:
    """Split **bold** and *italic* markdown spans."""
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out
